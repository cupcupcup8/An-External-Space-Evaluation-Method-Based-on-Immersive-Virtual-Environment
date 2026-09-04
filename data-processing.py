import pandas as pd
import datetime
import os
import numpy as np
import openpyxl

TIME_COLUMN_INDEX = 0


def is_valid_time(value):
    return value is not None and not (isinstance(value, float) and np.isnan(value))


def convert_time_to_seconds(value):
    if not is_valid_time(value):
        return np.nan
    try:
        if isinstance(value, datetime.time):
            return (value.hour * 3600 + value.minute * 60 +
                    value.second + value.microsecond / 1000000)
        elif isinstance(value, str):
            parts = value.split(':')
            minutes = int(parts[0])
            sec_parts = parts[1].split('.')
            seconds = int(sec_parts[0])
            ms = int(sec_parts[1]) if len(sec_parts) > 1 else 0
            return minutes * 60 + seconds + ms / 1000
        else:
            return float(value) * 86400  # Excel time stored in days
    except:
        return np.nan


def format_seconds(seconds):
    if not is_valid_time(seconds):
        return "invalid time"
    return f"{seconds:.3f}s"


def get_visible_rows(file_path, sheet_name):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb[sheet_name]
    visible = []
    for row in range(1, ws.max_row + 1):
        if not ws.row_dimensions[row].hidden:
            visible.append(ws.cell(row=row, column=TIME_COLUMN_INDEX + 1).value)
    wb.close()
    return visible


def analyze_time_series(time_series, sheet_name):
    total_points = len(time_series)

    seconds = []
    invalid_count = 0
    for value in time_series:
        sec = convert_time_to_seconds(value)
        if np.isnan(sec):
            invalid_count += 1
        seconds.append(sec)

    # Group consecutive points with gaps <= 1 second
    groups = []
    current = []
    for sec in seconds:
        if np.isnan(sec):
            continue
        if not current or abs(sec - current[-1]) <= 1.0:
            current.append(sec)
        else:
            groups.append((current[0], current[-1], len(current)))
            current = [sec]
    if current:
        groups.append((current[0], current[-1], len(current)))

    result = {
        "sheet_name": sheet_name,
        "visible_points": total_points,
        "invalid_values": invalid_count,
        "group_count": len(groups),
        "status": "done",
    }

    significant_count = 0
    total_duration = 0
    group_details = []
    first_significant = None

    for i, (start, end, count) in enumerate(groups):
        duration = abs(end - start)
        info = {
            "group_no": i + 1,
            "start_sec": start,
            "end_sec": end,
            "duration_sec": duration,
            "points": count,
            "isolated": count == 1,
        }
        group_details.append(info)

        if duration > 0.2:
            significant_count += 1
            total_duration += duration
            if first_significant is None:
                first_significant = info

    result["significant_groups"] = significant_count
    result["total_duration_sec"] = total_duration
    result["group_details"] = group_details

    if first_significant:
        result.update({
            "first_group_start_sec": first_significant["start_sec"],
            "first_group_end_sec": first_significant["end_sec"],
            "first_group_duration_sec": first_significant["duration_sec"],
        })
    else:
        result.update({
            "first_group_start_sec": np.nan,
            "first_group_end_sec": np.nan,
            "first_group_duration_sec": np.nan,
        })

    return result


def process_all_sheets(file_path, output_path):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    wb.close()

    all_results = []
    all_group_details = []

    for sheet_name in sheet_names:
        visible_data = get_visible_rows(file_path, sheet_name)
        sheet_result = analyze_time_series(visible_data, sheet_name)
        all_results.append(sheet_result)
        for group in sheet_result["group_details"]:
            group["sheet_name"] = sheet_name
            all_group_details.append(group)

    results_df = pd.DataFrame(all_results)
    group_details_df = pd.DataFrame(all_group_details)

    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        results_df.to_excel(writer, sheet_name='time_analysis_summary', index=False)
        if not group_details_df.empty:
            group_details_df.to_excel(writer, sheet_name='group_details', index=False)

    return {"summary": results_df, "group_details": group_details_df}


file_path = r"E:\pythonProject1\YDSJ\4-40\4-40-GOU_Processed.xlsx"
output_path = r"E:\pythonProject1\YDSJ\4-40\4-40-GOU_0904Processed2.xlsx"

TIME_COLUMN_INDEX = 0

process_all_sheets(file_path, output_path)
